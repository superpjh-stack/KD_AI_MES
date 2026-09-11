"""입고재고관리 005~008 (개발1).

입고 흐름 (goal.md §5 개발1 · SF-TD4-005)
  스마트패드로 자재 LOT·규격(재질·두께)·중량 등록 → **MTC(재질성적서) 첨부** · 검사항목 판정
  → 정합성 검증·표준화 → `INV_STOCKS` 반영 → `INV_SUPPLIER_QUALITY` 누적.

**ERP(이카운트) 연계는 사업계획서 안에서 상충한다**(D-07). 실 연동은 범위 밖이므로
**Excel 적재를 정식 입력 경로**로 두고 `IF_ERP_RECEIPTS` 에 성공·실패 이력을 남긴다.

`/inv/009`(입고 AI Agent)는 개발3 의 `routers/agt.py` 다 — 여기서 다루지 않는다.
"""
from __future__ import annotations

import io
from datetime import datetime
from typing import Any

from fastapi import APIRouter, Request

from .bas import (anchor_label, can_write, code_names, code_options, dt, guard,  # noqa: F401
                  lot_link, opt_num, paginate, project_link, redirect,
                  require_fields, screen_page, search_spec, undetermined, val)
from ..util import clock, codes, csrf, http

import conn                                              # noqa: E402  (bas 가 db 경로를 넣는다)

router = APIRouter()
SCREENS = ("MES-TD3-005", "MES-TD3-006", "MES-TD3-007", "MES-TD3-008")

# TD5 비고가 정한 열거값 — 코드 그룹이 아니라 컬럼 비고다(D-32 29건에 없다).
INSPECT_RESULTS = ("합격", "불합격", "보류")
INPUT_DEVICES = ("스마트패드", "현장POP", "Web")
SYNC_STATES = ("미전송", "전송", "오류")
HIST_TYPES = ("입고", "투입", "반출", "반입", "출하")

# Excel 적재 헤더 — `IF_ERP_RECEIPTS` + `INV_MATERIAL_LOTS` 컬럼에 1:1 로 대응한다.
EXCEL_COLUMNS = ("ERP전표번호", "품목코드", "재질", "두께mm", "공급처코드", "공급처명",
                 "입고수량", "중량kg", "MTC번호", "입고일시", "검사판정")


def _opts(values) -> list[dict[str, str]]:
    return [{"value": v, "label": v} for v in values]


def _suppliers() -> list[dict[str, str]]:
    rows = conn.q("select SUPPLIER_ID, SUPPLIER_CODE, SUPPLIER_NAME from INV_SUPPLIERS "
                  "where USE_YN = 'Y' order by SUPPLIER_NAME")
    return [{"value": str(r["supplier_id"]), "label": f"{r['supplier_name']} ({r['supplier_code']})"}
            for r in rows]


# ── 채번 (progress-dev1.md §1 공표 · BAS_COMMON_CODES 그룹 'LOT채번') ────────
def _numbering(key: str) -> str:
    row = conn.q1("select ATTR1 from BAS_COMMON_CODES where CODE_GROUP = 'LOT채번' and CODE_VALUE = %s",
                  (key,))
    if not row or not row["attr1"]:
        raise http.fail("internal",
                        f"채번 규칙 '{key}' 가 BAS_COMMON_CODES 그룹 'LOT채번' 에 없다 — "
                        "`uv run python db/seed_dev1.py`")
    return row["attr1"]


def _next_receipt_no(cur, when: datetime) -> str:
    """`RC-YYYY-NNNN` — 연도별 일련."""
    _numbering("RECEIPT_NO")                 # 규칙이 없으면 여기서 터진다(조용히 만들지 않는다)
    prefix = f"RC-{when.year}-"
    cur.execute("select RECEIPT_NO from INV_RECEIPTS where RECEIPT_NO like %s "
                "order by RECEIPT_NO desc limit 1", (prefix + "%",))
    last = cur.fetchone()
    n = int(last["receipt_no"].rsplit("-", 1)[-1]) + 1 if last else 1
    return f"{prefix}{n:04d}"


def _next_lot_no(cur, when: datetime) -> str:
    """`LOT-YYMMDD-NN` — 일자별 일련."""
    _numbering("MATERIAL_LOT_NO")
    prefix = f"LOT-{when:%y%m%d}-"
    cur.execute("select LOT_NO from INV_MATERIAL_LOTS where LOT_NO like %s "
                "order by LOT_NO desc limit 1", (prefix + "%",))
    last = cur.fetchone()
    n = int(last["lot_no"].rsplit("-", 1)[-1]) + 1 if last else 1
    return f"{prefix}{n:02d}"


def _receive(cur, *, supplier_id: int, item_code: str, material: str | None,
             thickness: float | None, weight: float | None, qty: float,
             mtc_no: str | None, inspect: str | None, when: datetime,
             device: str, user_id: int | None, project_id: int | None = None) -> tuple[str, str]:
    """입고 1건 = LOT 생성 → 입고 기록 → 재고 반영 → 이력 적재. **한 트랜잭션이다.**

    코드성 FK 검증은 **호출 전에** 끝나 있어야 한다(D-32).
    """
    lot_no = _next_lot_no(cur, when)
    receipt_no = _next_receipt_no(cur, when)

    cur.execute(
        "insert into INV_MATERIAL_LOTS "
        "(LOT_NO, ITEM_CODE, MATERIAL, THICKNESS_MM, WEIGHT_KG, CURRENT_QTY, LOT_STATUS, CREATED_DT) "
        "values (%s,%s,%s,%s,%s,%s,'입고', now()) returning LOT_ID",
        (lot_no, item_code, material, thickness, weight, qty))
    lot_id = cur.fetchone()["lot_id"]

    cur.execute(
        "insert into INV_RECEIPTS "
        "(RECEIPT_NO, PROJECT_ID, SUPPLIER_ID, LOT_ID, RECEIPT_DT, RECEIPT_QTY, WEIGHT_KG, "
        " MTC_NO, INSPECT_RESULT, INPUT_DEVICE, ERP_SYNC_STATUS, CREATED_BY, CREATED_DT) "
        "values (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,'미전송',%s, now())",
        (receipt_no, project_id, supplier_id, lot_id, when, qty, weight,
         mtc_no, inspect, device, user_id))

    # 재고 반영 — 합격만 가용 수량에 들어간다. 보류·불합격은 현재고에만 잡힌다.
    available = qty if inspect == "합격" else 0
    cur.execute(
        "insert into INV_STOCKS (LOT_ID, STOCK_QTY, AVAILABLE_QTY, LAST_IN_DT, CREATED_DT) "
        "values (%s,%s,%s,%s, now())", (lot_id, qty, available, when))

    # 원자재 이력 — 런타임 누적표다(D-102). 입고 시점이 첫 행이다.
    cur.execute(
        "insert into INV_MATERIAL_HISTORY "
        "(LOT_ID, HIST_TYPE, EVENT_DT, PROCESS_CODE, EVENT_QTY, USER_ID, REMARK, CREATED_DT) "
        "values (%s,'입고',%s,%s,%s,%s,%s, now())",
        (lot_id, when, "P20", qty, user_id,          # P20 = 입고검사 (공통 시드 그룹 '공정')
         f"입고 {receipt_no} · 판정 {inspect or '미판정'}"))
    return receipt_no, lot_no


# ═════════════════════════════════════════════════════════════════════════
# 005 입고관리
# ═════════════════════════════════════════════════════════════════════════
@router.get("/inv/005")
async def receipt(request: Request):
    sid = "MES-TD3-005"
    scr, td3 = guard(request, sid)
    q = request.query_params
    where, params = ["1=1"], []
    if q.get("day"):
        where.append("r.RECEIPT_DT::date = %s"); params.append(q["day"])
    if q.get("no"):
        where.append("r.RECEIPT_NO ilike %s"); params.append(f"%{q['no']}%")
    if q.get("sup"):
        where.append("r.SUPPLIER_ID = %s"); params.append(q["sup"])
    if q.get("item"):
        where.append("l.ITEM_CODE = %s"); params.append(q["item"])
    if q.get("lot"):
        where.append("l.LOT_NO ilike %s"); params.append(f"%{q['lot']}%")
    if q.get("judge"):
        where.append("r.INSPECT_RESULT = %s"); params.append(q["judge"])
    cond = " and ".join(where)
    join = ("from INV_RECEIPTS r "
            "join INV_MATERIAL_LOTS l on l.LOT_ID = r.LOT_ID "
            "join INV_SUPPLIERS s on s.SUPPLIER_ID = r.SUPPLIER_ID")

    total = int(conn.q1(f"select count(*) as n {join} where {cond}", params)["n"])
    pager, size, off = paginate(request, total)
    rows = conn.q(
        f"select r.RECEIPT_NO, r.RECEIPT_DT, s.SUPPLIER_NAME, l.LOT_NO, r.RECEIPT_QTY, "
        f"r.INSPECT_RESULT {join} where {cond} order by r.RECEIPT_DT desc, r.RECEIPT_NO desc "
        f"limit %s offset %s", [*params, size, off])
    grid = [[str(off + i), r["receipt_no"], dt(r["receipt_dt"]), r["supplier_name"],
             lot_link(r["lot_no"]), f"{r['receipt_qty']:,.3f}",
             val(r["inspect_result"], "D-101")]
            for i, r in enumerate(rows, 1)]

    judged = conn.q1(
        "select count(*) filter (where INSPECT_RESULT = '합격') as ok, "
        "count(*) filter (where INSPECT_RESULT = '불합격') as ng, "
        "count(*) filter (where INSPECT_RESULT is null) as none from INV_RECEIPTS")
    items, materials = code_options("품목"), code_options("재질")
    blocked = not items

    return screen_page(
        request, sid, _loaded=(scr, td3),
        search=search_spec(td3, request, [
            ("day", "date"), ("no", "text"), ("sup", "text", _suppliers()),
            ("item", "text", items), ("lot", "text"), ("judge", "text", _opts(INSPECT_RESULTS))]),
        rows=grid, total=total, pager=pager, grid_title="입고",
        empty_notice=http.not_collected("D-101") +
        " — 현행 입고 데이터는 Excel 약 1천개와 종이 MTC 로 관리되어 초기 이관 범위를 착수 시 "
        "확정해야 한다(SF-TD3-005 체크). 아래 등록 폼 또는 007 Excel 적재로 들어온다.",
        cards=[{"label": "입고", "value": f"{total} 건"},
               {"label": "합격", "value": f"{int(judged['ok'])} 건"},
               {"label": "불합격", "value": f"{int(judged['ng'])} 건"},
               {"label": "미판정", "value": f"{int(judged['none'])} 건"}],
        notices=([{"kind": "bad",
                   "text": "공통코드 그룹 '품목' 이 비어 있어 입고를 등록할 수 없다 — "
                           "032 코드관리에서 품목·재질 코드를 먼저 등록한다 (D-47 · D-101)"}]
                 if blocked else []) +
                [{"kind": "notice",
                  "text": "자재 LOT 을 누르면 024 공정이력조회로 간다 (G-08)"}],
        form={
            "title": "입고 등록 (스마트패드 · 관리자 Web)", "action": "/inv/005", "submit": "저장",
            "note": "LOT 번호·입고번호는 채번 규칙(BAS_COMMON_CODES 그룹 'LOT채번')으로 자동 부여된다. "
                    "품목코드·재질은 공통코드에 있는 값만 저장된다(D-32 — 어기면 422). "
                    "MTC(재질성적서) 번호는 입고검사 근거다. 저장은 LOT·입고·재고·이력을 한 트랜잭션으로 쓴다.",
            "fields": [
                {"label": "공급처", "name": "supplier_id", "required": True, "options": _suppliers()},
                {"label": "품목코드", "name": "item_code", "required": True, "options": items},
                {"label": "재질", "name": "material", "options": materials},
                {"label": "두께(mm)", "name": "thickness_mm", "type": "number", "step": "any"},
                {"label": "입고수량", "name": "receipt_qty", "type": "number", "step": "any",
                 "required": True},
                {"label": "중량(kg)", "name": "weight_kg", "type": "number", "step": "any"},
                {"label": "MTC 번호", "name": "mtc_no"},
                {"label": "검사 판정", "name": "inspect_result", "options": _opts(INSPECT_RESULTS)},
                {"label": "입력 단말", "name": "input_device", "options": _opts(INPUT_DEVICES)},
                {"label": "입고일시", "name": "receipt_dt", "type": "datetime-local"},
            ],
        },
        disabled_note="등록·저장은 아래 폼이 담당한다. 삭제·엑셀은 미구현이다 — "
                      "조용히 성공한 척하지 않는다(G-30).",
    )


@router.post("/inv/005")
async def receipt_save(request: Request):
    sid = "MES-TD3-005"
    f = await request.form()
    csrf.require(request, f.get("_csrf"))
    guard(request, sid, write=True)
    v = require_fields(f, ("supplier_id", "item_code", "receipt_qty"))
    material = (f.get("material") or "").strip() or None
    inspect = (f.get("inspect_result") or "").strip() or None
    device = (f.get("input_device") or "").strip() or "Web"

    # D-32 — DB 가 막지 않는다. 저장 전에 막는다.
    codes.require_code("INV_MATERIAL_LOTS", "ITEM_CODE", v["item_code"])
    codes.require_code("INV_MATERIAL_LOTS", "MATERIAL", material, allow_empty=True)
    if inspect and inspect not in INSPECT_RESULTS:
        raise http.fail("validation", f"검사 판정은 {INSPECT_RESULTS} 중 하나다: {inspect!r}")
    if device not in INPUT_DEVICES:
        raise http.fail("validation", f"입력 단말은 {INPUT_DEVICES} 중 하나다: {device!r}")

    qty = opt_num(f, "receipt_qty")
    if qty is None or qty <= 0:
        raise http.fail("validation", "입고수량은 0보다 커야 한다")
    if not conn.q1("select 1 as x from INV_SUPPLIERS where SUPPLIER_ID = %s and USE_YN='Y'",
                   (v["supplier_id"],)):
        raise http.fail("validation", f"공급처 {v['supplier_id']} 를 찾을 수 없다")

    raw = (f.get("receipt_dt") or "").strip()
    # 시간 기준일은 앵커다 — `datetime.now()` 를 기본값으로 쓰지 않는다(§10-3).
    when = datetime.fromisoformat(raw) if raw else clock.anchor()

    with conn.tx() as cur:
        _receive(cur, supplier_id=int(v["supplier_id"]), item_code=v["item_code"],
                 material=material, thickness=opt_num(f, "thickness_mm"),
                 weight=opt_num(f, "weight_kg"), qty=qty,
                 mtc_no=(f.get("mtc_no") or "").strip() or None, inspect=inspect,
                 when=when, device=device,
                 user_id=getattr(getattr(request.state, "session", None), "user_id", None))
    return redirect("/inv/005")


# ═════════════════════════════════════════════════════════════════════════
# 006 원자재 이력조회
# ═════════════════════════════════════════════════════════════════════════
@router.get("/inv/006")
async def material_history(request: Request):
    sid = "MES-TD3-006"
    scr, td3 = guard(request, sid)
    q = request.query_params
    where, params = ["1=1"], []
    if q.get("lot"):
        where.append("l.LOT_NO ilike %s"); params.append(f"%{q['lot']}%")
    if q.get("item"):
        where.append("l.ITEM_CODE = %s"); params.append(q["item"])
    if q.get("from"):
        where.append("h.EVENT_DT >= %s"); params.append(q["from"])
    if q.get("to"):
        where.append("h.EVENT_DT < (%s::date + 1)"); params.append(q["to"])
    if q.get("kind"):
        where.append("h.HIST_TYPE = %s"); params.append(q["kind"])
    if q.get("proc"):
        where.append("h.PROCESS_CODE = %s"); params.append(q["proc"])
    if q.get("plot"):
        where.append("h.PRODUCT_LOT_NO ilike %s"); params.append(f"%{q['plot']}%")
    cond = " and ".join(where)
    join = "from INV_MATERIAL_HISTORY h join INV_MATERIAL_LOTS l on l.LOT_ID = h.LOT_ID"

    total = int(conn.q1(f"select count(*) as n {join} where {cond}", params)["n"])
    pager, size, off = paginate(request, total)
    rows = conn.q(
        f"select l.LOT_NO, h.HIST_TYPE, h.EVENT_DT, h.PROCESS_CODE, h.EVENT_QTY, "
        f"h.PRODUCT_LOT_NO {join} where {cond} order by h.EVENT_DT desc, h.HIST_ID desc "
        f"limit %s offset %s", [*params, size, off])
    proc = code_names("공정")
    grid = [[str(off + i), lot_link(r["lot_no"]), r["hist_type"], dt(r["event_dt"]),
             (proc.get(r["process_code"], r["process_code"]) if r["process_code"]
              else undetermined("D-102")),
             f"{r['event_qty']:,.3f}",
             lot_link(r["product_lot_no"], "D-102") if r["product_lot_no"] else "-"]
            for i, r in enumerate(rows, 1)]

    linked = int(conn.q1("select count(*) as n from INV_MATERIAL_HISTORY "
                         "where PRODUCT_LOT_NO is not null")["n"])
    rate = f"{linked / total * 100:.1f} %" if total else "0.0 %"

    return screen_page(
        request, sid, _loaded=(scr, td3),
        search=search_spec(td3, request, [
            ("lot", "text"), ("item", "text", code_options("품목")), ("from", "date"),
            ("kind", "text", _opts(HIST_TYPES)), ("proc", "text", code_options("공정")),
            ("plot", "text")]),
        rows=grid, total=total, pager=pager, grid_title="자재 이력",
        empty_notice=http.not_collected("D-102") +
        " — 원자재 이력은 런타임 누적표다. 입고·투입·반출·반입·출하 시점에 쌓인다.",
        cards=[{"label": "이력", "value": f"{total} 건"},
               {"label": "제품 LOT 연결", "value": f"{linked} 건"},
               {"label": "연결률", "value": rate}],
        notices=[{"kind": "notice",
                  "text": "LOT 기반 데이터 매핑 성공률 목표 85% 이상 (사업계획서 2.7.4 · D-23) — "
                          "위 연결률은 실측이다"},
                 {"kind": "notice",
                  "text": "자재 LOT·제품 LOT 을 누르면 024 공정이력조회로 간다 (G-08)"}],
        disabled_note="이력추적은 LOT 링크(→ 024 공정이력조회)가 담당한다. 엑셀은 미구현이다(G-30).",
    )


# ═════════════════════════════════════════════════════════════════════════
# 007 입고 데이터관리 — 정합성 검증·표준화 + ERP 연계 (D-07)
# ═════════════════════════════════════════════════════════════════════════
def _excel_rows(raw: bytes) -> list[dict[str, Any]]:
    """업로드된 .xlsx 를 행 dict 목록으로. 헤더가 다르면 **422** — 조용히 넘기지 않는다."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise http.fail("validation", "빈 파일이다")
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    missing = [c for c in EXCEL_COLUMNS if c not in header]
    if missing:
        raise http.fail("validation",
                        f"Excel 헤더 누락: {missing} — 필요한 열은 {list(EXCEL_COLUMNS)}")
    idx = {c: header.index(c) for c in EXCEL_COLUMNS}
    out = []
    for r in rows[1:]:
        if r is None or all(c is None or str(c).strip() == "" for c in r):
            continue
        out.append({c: (r[idx[c]] if idx[c] < len(r) else None) for c in EXCEL_COLUMNS})
    return out


def _as_dt(value: Any) -> datetime:
    if isinstance(value, datetime):
        return value
    if value is None or str(value).strip() == "":
        return clock.anchor()             # 기준일은 앵커다 (§10-3)
    return datetime.fromisoformat(str(value).strip())


def _as_num(value: Any, field: str) -> float | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        raise http.fail("validation", f"{field} 은 숫자다: {value!r}") from None


@router.get("/inv/007")
async def receipt_data(request: Request):
    sid = "MES-TD3-007"
    scr, td3 = guard(request, sid)
    q = request.query_params
    where, params = ["1=1"], []
    if q.get("from"):
        where.append("r.RECEIPT_DT >= %s"); params.append(q["from"])
    if q.get("to"):
        where.append("r.RECEIPT_DT < (%s::date + 1)"); params.append(q["to"])
    if q.get("no"):
        where.append("r.RECEIPT_NO ilike %s"); params.append(f"%{q['no']}%")
    if q.get("erp"):
        where.append("e.ERP_DOC_NO ilike %s"); params.append(f"%{q['erp']}%")
    if q.get("sync"):
        where.append("r.ERP_SYNC_STATUS = %s"); params.append(q["sync"])
    cond = " and ".join(where)
    # ERP 연계 이력은 품목·수량·일자로 맞춘다 — 전표번호를 입고에 들고 있지 않기 때문이다(D-07).
    join = ("from INV_RECEIPTS r "
            "join INV_MATERIAL_LOTS l on l.LOT_ID = r.LOT_ID "
            "left join lateral (select ERP_DOC_NO, IF_STATUS, ERROR_MSG, IF_DT "
            "                   from IF_ERP_RECEIPTS x "
            "                   where x.ITEM_CODE = l.ITEM_CODE and x.RECEIPT_QTY = r.RECEIPT_QTY "
            "                     and x.RECEIPT_DT = r.RECEIPT_DT "
            "                   order by x.IF_DT desc limit 1) e on true")

    total = int(conn.q1(f"select count(*) as n {join} where {cond}", params)["n"])
    pager, size, off = paginate(request, total)
    rows = conn.q(
        f"select r.RECEIPT_NO, r.RECEIPT_QTY, r.ERP_SYNC_STATUS, r.INSPECT_RESULT, r.MTC_NO, "
        f"e.ERP_DOC_NO, e.IF_STATUS, e.ERROR_MSG, e.IF_DT {join} where {cond} "
        f"order by r.RECEIPT_DT desc limit %s offset %s", [*params, size, off])

    def verdict(r) -> Any:
        """검증 결과 — 실측이다. 판정 근거가 없으면 미확정을 낸다."""
        if r["error_msg"]:
            return {"text": r["error_msg"][:60], "notice": True}
        if r["inspect_result"] is None:
            return {"text": "검사 판정 없음", "notice": True}
        if not r["mtc_no"]:
            return {"text": "MTC 미첨부", "notice": True}
        return "정합"

    grid = [[str(off + i), r["receipt_no"], r["erp_doc_no"] or "-",
             f"{r['receipt_qty']:,.3f}", val(r["erp_sync_status"], "D-07"),
             verdict(r), dt(r["if_dt"], "D-07")]
            for i, r in enumerate(rows, 1)]

    checks = conn.q(
        "select CHECK_AXIS, TARGET_DESC, TOTAL_CNT, VALID_CNT, ACHIEVE_RATE, JUDGE_RESULT, CHECKED_DT "
        "from DAT_QUALITY_CHECKS order by CHECKED_DT desc limit 20")
    check_rows = [[str(i), c["check_axis"], c["target_desc"], str(c["total_cnt"]),
                   str(c["valid_cnt"]), val(c["achieve_rate"], "D-09"),
                   val(c["judge_result"], "D-09"), dt(c["checked_dt"])]
                  for i, c in enumerate(checks, 1)]
    if_total = int(conn.q1("select count(*) as n from IF_ERP_RECEIPTS")["n"])
    if_fail = int(conn.q1("select count(*) as n from IF_ERP_RECEIPTS where IF_STATUS='실패'")["n"])

    return screen_page(
        request, sid, _loaded=(scr, td3),
        search=search_spec(td3, request, [
            ("from", "date"), ("no", "text"), ("erp", "text"),
            ("sync", "text", _opts(SYNC_STATES)), ("verify", "text")]),
        rows=grid, total=total, pager=pager, grid_title="입고 연계",
        empty_notice=http.not_collected("D-07") + " — 입고가 0건이다. 아래 Excel 적재가 정식 입력 경로다.",
        cards=[{"label": "입고", "value": f"{total} 건"},
               {"label": "ERP 연계 이력", "value": f"{if_total} 건"},
               {"label": "연계 실패", "value": f"{if_fail} 건"},
               {"label": "품질검증", "value": f"{len(checks)} 건"}],
        notices=[{"kind": "bad",
                  "text": "ERP(이카운트) 보유·연계 범위가 사업계획서 안에서 상충한다 (D-07) — "
                          "실 연동은 범위 밖이다. Excel 적재를 정식 입력으로 두고 "
                          "IF_ERP_RECEIPTS 에 이력을 남긴다"}],
        form={
            "title": "Excel 적재 (정식 입력 경로 — D-07)", "action": "/inv/007",
            "submit": "적재", "upload": True,
            "note": f"헤더: {' · '.join(EXCEL_COLUMNS)}. "
                    "행마다 품목코드·재질을 공통코드로 검증하고(D-32), 성공·실패를 "
                    "IF_ERP_RECEIPTS 에 남긴다. 실패 행은 적재하지 않는다 — 조용히 넘기지 않는다(G-30).",
            "fields": [{"label": "Excel 파일(.xlsx)", "name": "file", "type": "file", "required": True}],
        },
        panels=[{
            "title": "데이터 품질 검증 (DAT_QUALITY_CHECKS)",
            "columns": ["No", "검증 축", "대상 데이터", "전체", "정상", "달성률(%)", "판정", "검증 일시"],
            "rows": check_rows,
            "empty_notice": http.not_collected("D-09") +
            " — 검증을 아직 실행하지 않았다. 아래 '정합성 검증' 이 실측으로 기록한다.",
            "note": "목표 85% 이상 (사업계획서 2.7.4). 달성률은 실측이고 시드로 채우지 않는다.",
        }],
        disabled_note="'재전송'·'표준화' 는 아래 버튼으로 실행한다. 엑셀 다운로드는 미구현이다(G-30).",
        extra_actions=[
            {"label": "정합성 검증", "action": "/inv/007", "name": "action", "value": "verify"},
            {"label": "ERP 재전송", "action": "/inv/007", "name": "action", "value": "resend"},
        ],
    )


@router.post("/inv/007")
async def receipt_data_action(request: Request):
    sid = "MES-TD3-007"
    form = await request.form()
    csrf.require(request, form.get("_csrf"))
    guard(request, sid, write=True)
    action = (form.get("action") or "upload").strip()
    user_id = getattr(getattr(request.state, "session", None), "user_id", None)

    if action == "verify":
        _verify_receipts()
        return redirect("/inv/007")
    if action == "resend":
        _resend_receipts()
        return redirect("/inv/007")
    if action != "upload":
        raise http.fail("validation", f"알 수 없는 동작: {action!r}")

    upload = form.get("file")
    raw = await upload.read() if hasattr(upload, "read") else None
    if not raw:
        raise http.fail("validation", "Excel 파일이 없다")
    rows = _excel_rows(raw)
    if not rows:
        raise http.fail("validation", "적재할 행이 없다")

    ok = fail = 0
    for i, r in enumerate(rows, 2):        # 2행부터가 데이터다
        item = str(r["품목코드"] or "").strip()
        material = str(r["재질"] or "").strip() or None
        sup_code = str(r["공급처코드"] or "").strip()
        sup_name = str(r["공급처명"] or "").strip()
        erp_no = str(r["ERP전표번호"] or "").strip()
        judge = str(r["검사판정"] or "").strip() or None
        when = _as_dt(r["입고일시"])
        qty = _as_num(r["입고수량"], f"{i}행 입고수량")

        problem = None
        if not item:
            problem = "품목코드 누락"
        elif not codes.validate_code("품목", item):
            problem = f"품목코드 '{item}' 가 공통코드 그룹 '품목' 에 없다 (D-32)"
        elif material and not codes.validate_code("재질", material):
            problem = f"재질 '{material}' 가 공통코드 그룹 '재질' 에 없다 (D-32)"
        elif not sup_code:
            problem = "공급처코드 누락"
        elif qty is None or qty <= 0:
            problem = "입고수량이 0 이하"
        elif judge and judge not in INSPECT_RESULTS:
            problem = f"검사판정 '{judge}' 는 {INSPECT_RESULTS} 가 아니다"

        if problem:
            fail += 1
            conn.x(
                "insert into IF_ERP_RECEIPTS "
                "(ERP_DOC_NO, ITEM_CODE, SUPPLIER_CODE, RECEIPT_QTY, RECEIPT_DT, IF_DIRECTION, "
                " IF_STATUS, ERROR_MSG, IF_DT, CREATED_DT) "
                "values (%s,%s,%s,%s,%s,'수신','실패',%s,%s, now())",
                (erp_no or f"(전표번호 없음 {i}행)", item or "(없음)", sup_code or None,
                 qty or 0, when, f"{i}행: {problem}", when))
            continue

        with conn.tx() as cur:
            cur.execute("select SUPPLIER_ID from INV_SUPPLIERS where SUPPLIER_CODE = %s", (sup_code,))
            found = cur.fetchone()
            if found:
                supplier_id = found["supplier_id"]
            else:
                # 공급처 마스터 등록 화면이 45화면에 없다 — Excel 적재가 유일한 유입 경로다(D-101).
                cur.execute(
                    "insert into INV_SUPPLIERS "
                    "(SUPPLIER_CODE, SUPPLIER_NAME, RISK_YN, USE_YN, CREATED_DT) "
                    "values (%s,%s,'N','Y', now()) returning SUPPLIER_ID",
                    (sup_code, sup_name or sup_code))
                supplier_id = cur.fetchone()["supplier_id"]

            _receive(cur, supplier_id=supplier_id, item_code=item, material=material,
                     thickness=_as_num(r["두께mm"], f"{i}행 두께"),
                     weight=_as_num(r["중량kg"], f"{i}행 중량"), qty=qty,
                     mtc_no=str(r["MTC번호"] or "").strip() or None, inspect=judge,
                     when=when, device="Web", user_id=user_id)
            cur.execute(
                "insert into IF_ERP_RECEIPTS "
                "(ERP_DOC_NO, ITEM_CODE, SUPPLIER_CODE, RECEIPT_QTY, RECEIPT_DT, IF_DIRECTION, "
                " IF_STATUS, IF_DT, CREATED_DT) values (%s,%s,%s,%s,%s,'수신','성공',%s, now())",
                (erp_no or f"XL-{when:%Y%m%d}-{i:04d}", item, sup_code, qty, when, when))
        ok += 1

    _verify_receipts()      # 적재 직후 정합성을 실측해 기록한다
    return redirect(f"/inv/007?loaded={ok}&failed={fail}")


def _verify_receipts() -> None:
    """정합성 검증 — **실측이다.** 목표 85%(2.7.4). 값을 지어내지 않는다.

    정확성: 검사 판정이 있고 MTC 가 첨부된 입고 비율.
    연계성: 자재 LOT 이 제품 LOT 으로 이어진 이력 비율 (G-08 · D-23).
    """
    when = clock.anchor()
    total = int(conn.q1("select count(*) as n from INV_RECEIPTS")["n"])
    valid = int(conn.q1("select count(*) as n from INV_RECEIPTS "
                        "where INSPECT_RESULT is not null and MTC_NO is not null "
                        "and MTC_NO <> ''")["n"])
    hist = int(conn.q1("select count(*) as n from INV_MATERIAL_HISTORY")["n"])
    linked = int(conn.q1("select count(*) as n from INV_MATERIAL_HISTORY "
                         "where PRODUCT_LOT_NO is not null")["n"])
    for axis, desc, tot, ok in (
            ("정확성", "입고 — 검사 판정·MTC(재질성적서) 첨부", total, valid),
            ("연계성", "자재 LOT → 제품 LOT 이력 연결 (G-08)", hist, linked)):
        rate = round(ok / tot * 100, 3) if tot else None
        conn.x(
            "insert into DAT_QUALITY_CHECKS "
            "(CHECK_AXIS, TARGET_DESC, TOTAL_CNT, VALID_CNT, ACHIEVE_RATE, JUDGE_RESULT, "
            " CHECKED_DT, ACTION_COMMENT, CREATED_DT) values (%s,%s,%s,%s,%s,%s,%s,%s, now())",
            (axis, desc, tot, ok, rate,
             None if rate is None else ("충족" if rate >= 85 else "미충족"), when,
             "전체 0건 — 달성률 판정 불가" if tot == 0 else None))


def _resend_receipts() -> None:
    """`ERP_SYNC_STATUS` 가 '미전송'·'오류' 인 입고를 다시 연계 대기로 남긴다.

    **실 ERP 로 보내지 않는다**(D-07 — 연동은 범위 밖). 보낸 척하지 않고 '대기' 로 기록한다.
    """
    when = clock.anchor()
    rows = conn.q(
        "select r.RECEIPT_NO, r.RECEIPT_QTY, r.RECEIPT_DT, l.ITEM_CODE, s.SUPPLIER_CODE "
        "from INV_RECEIPTS r join INV_MATERIAL_LOTS l on l.LOT_ID = r.LOT_ID "
        "join INV_SUPPLIERS s on s.SUPPLIER_ID = r.SUPPLIER_ID "
        "where r.ERP_SYNC_STATUS in ('미전송','오류')")
    for r in rows:
        conn.x(
            "insert into IF_ERP_RECEIPTS "
            "(ERP_DOC_NO, ITEM_CODE, SUPPLIER_CODE, RECEIPT_QTY, RECEIPT_DT, IF_DIRECTION, "
            " IF_STATUS, ERROR_MSG, IF_DT, CREATED_DT) "
            "values (%s,%s,%s,%s,%s,'송신','대기',%s,%s, now())",
            (r["receipt_no"], r["item_code"], r["supplier_code"], r["receipt_qty"],
             r["receipt_dt"], "ERP 실 연동은 범위 밖이다 (D-07) — 연계 대기로만 기록한다", when))


# ═════════════════════════════════════════════════════════════════════════
# 008 공급처 품질분석
# ═════════════════════════════════════════════════════════════════════════
def _period(when: datetime) -> str:
    """평가 기간 — TD5 `EVAL_PERIOD` 비고 'YYYY-MM 또는 분기'. 분기로 집계한다."""
    return f"{when.year}-Q{(when.month - 1) // 3 + 1}"


@router.get("/inv/008")
async def supplier_quality(request: Request):
    sid = "MES-TD3-008"
    scr, td3 = guard(request, sid)
    q = request.query_params
    where, params = ["1=1"], []
    if q.get("period"):
        where.append("sq.EVAL_PERIOD = %s"); params.append(q["period"])
    if q.get("sup"):
        where.append("sq.SUPPLIER_ID = %s"); params.append(q["sup"])
    if q.get("type"):
        where.append("s.SUPPLY_TYPE = %s"); params.append(q["type"])
    if q.get("grade"):
        where.append("sq.GRADE = %s"); params.append(q["grade"])
    if q.get("risk"):
        where.append("s.RISK_YN = %s"); params.append(q["risk"])
    cond = " and ".join(where)
    join = "from INV_SUPPLIER_QUALITY sq join INV_SUPPLIERS s on s.SUPPLIER_ID = sq.SUPPLIER_ID"

    total = int(conn.q1(f"select count(*) as n {join} where {cond}", params)["n"])
    pager, size, off = paginate(request, total)
    rows = conn.q(
        f"select s.SUPPLIER_NAME, sq.EVAL_PERIOD, sq.DELIVERY_CNT, sq.DEFECT_RATE, "
        f"sq.OTD_RATE, sq.GRADE {join} where {cond} "
        f"order by sq.EVAL_PERIOD desc, s.SUPPLIER_NAME limit %s offset %s", [*params, size, off])
    grid = [[str(off + i), r["supplier_name"], r["eval_period"], str(r["delivery_cnt"]),
             val(r["defect_rate"], "D-107"), val(r["otd_rate"], "D-107"),
             val(r["grade"], "D-107")]
            for i, r in enumerate(rows, 1)]

    # 공급처 마스터 — 담당자 정보는 **마스킹**해서 보여준다(G-29).
    from ..util import mask
    sup_rows = [[str(i), s["supplier_code"], s["supplier_name"],
                 val(s["supply_type"], "D-101"),
                 mask(s["contact_name"], "name") or undetermined("D-101"),
                 mask(s["contact_phone"], "phone") or undetermined("D-101"),
                 s["risk_yn"]]
                for i, s in enumerate(conn.q(
                    "select SUPPLIER_CODE, SUPPLIER_NAME, SUPPLY_TYPE, CONTACT_NAME, "
                    "CONTACT_PHONE, RISK_YN from INV_SUPPLIERS order by SUPPLIER_NAME"), 1)]
    periods = [{"value": r["eval_period"], "label": r["eval_period"]} for r in conn.q(
        "select distinct EVAL_PERIOD from INV_SUPPLIER_QUALITY order by EVAL_PERIOD desc")]
    grades = [{"value": r["grade"], "label": r["grade"]} for r in conn.q(
        "select distinct GRADE from INV_SUPPLIER_QUALITY where GRADE is not null order by GRADE")]

    return screen_page(
        request, sid, _loaded=(scr, td3),
        search=search_spec(td3, request, [
            ("period", "text", periods), ("sup", "text", _suppliers()),
            ("type", "text", code_options("공급구분")), ("grade", "text", grades),
            ("risk", "text", _opts(("Y", "N")))]),
        rows=grid, total=total, pager=pager, grid_title="공급처 평가",
        empty_notice=http.not_collected("D-103") +
        " — 공급처 품질평가는 입고·검사 누적 집계다. 아래 '평가산출' 이 실측으로 만든다.",
        cards=[{"label": "평가", "value": f"{total} 건"},
               {"label": "공급처", "value": f"{len(sup_rows)} 개사"},
               {"label": "평가 기간", "value": f"{len(periods)} 기"}],
        notices=[{"kind": "undetermined",
                  "text": "품질 편차·납기 준수율·종합 점수·등급은 산식과 원천 데이터가 정본에 없다 "
                          "— 산출하지 않고 미확정으로 둔다 (D-107)"},
                 {"kind": "notice",
                  "text": "불량률 = 불합격 ÷ 납품 × 100 (TD5 DEFECT_RATE 비고) — 이것만 실측한다"},
                 {"kind": "notice",
                  "text": "공급처 담당자명·연락처는 마스킹해서 표시한다 (G-29)"}],
        panels=[{
            "title": "공급처 마스터 (담당자 정보 마스킹 — G-29)",
            "columns": ["No", "공급처 코드", "공급처명", "공급 구분", "담당자", "연락처", "리스크"],
            "rows": sup_rows,
            "empty_notice": http.not_collected("D-101") +
            " — 공급처 목록이 정본에 없고 45화면에 마스터 등록 화면도 없다. "
            "007 Excel 적재가 유일한 유입 경로다.",
        }],
        disabled_note="'평가산출' 은 아래 버튼으로 실행한다. 리포트·엑셀은 미구현이다(G-30).",
        extra_actions=[{"label": "평가산출", "action": "/inv/008", "name": "action", "value": "evaluate"}],
    )


@router.post("/inv/008")
async def supplier_quality_evaluate(request: Request):
    sid = "MES-TD3-008"
    form = await request.form()
    csrf.require(request, form.get("_csrf"))
    guard(request, sid, write=True)
    if (form.get("action") or "evaluate").strip() != "evaluate":
        raise http.fail("validation", "알 수 없는 동작")

    rows = conn.q(
        "select r.SUPPLIER_ID, r.RECEIPT_DT, r.INSPECT_RESULT from INV_RECEIPTS r")
    agg: dict[tuple[int, str], list[int]] = {}
    for r in rows:
        key = (int(r["supplier_id"]), _period(r["receipt_dt"]))
        cell = agg.setdefault(key, [0, 0])
        cell[0] += 1
        if r["inspect_result"] == "불합격":
            cell[1] += 1

    for (supplier_id, period), (delivered, rejected) in agg.items():
        rate = round(rejected / delivered * 100, 3) if delivered else None
        existing = conn.q1(
            "select SQ_ID from INV_SUPPLIER_QUALITY where SUPPLIER_ID = %s and EVAL_PERIOD = %s",
            (supplier_id, period))
        if existing:
            conn.x(
                "update INV_SUPPLIER_QUALITY set DELIVERY_CNT=%s, REJECT_CNT=%s, DEFECT_RATE=%s, "
                "UPDATED_DT = now() where SQ_ID = %s",
                (delivered, rejected, rate, existing["sq_id"]))
        else:
            # 품질 편차·납기 준수율·종합 점수·등급은 산식이 정본에 없다 → 넣지 않는다(D-107).
            conn.x(
                "insert into INV_SUPPLIER_QUALITY "
                "(SUPPLIER_ID, EVAL_PERIOD, DELIVERY_CNT, REJECT_CNT, DEFECT_RATE, CREATED_DT) "
                "values (%s,%s,%s,%s,%s, now())",
                (supplier_id, period, delivered, rejected, rate))
    return redirect("/inv/008")
